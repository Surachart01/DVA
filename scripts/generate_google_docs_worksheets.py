#!/usr/bin/env python3
"""Generate Google Docs / Word editable worksheets (.docx) for Week 2."""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "TH Sarabun New"

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right.
    Value: dict(val='single', sz='6', space='0', color='CBD5E1')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", "4")}" w:space="{edge_data.get("space", "0")}" w:color="{edge_data.get("color", "CBD5E1")}"/>'
            tcBorders.append(parse_xml(tag))
        else:
            tag = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

def apply_run_font(run, font_name=FONT_NAME, size_pt=16, bold=False, italic=False, color_rgb=(15, 23, 42)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)

def add_styled_paragraph(doc, text="", font_name=FONT_NAME, size_pt=16, bold=False, italic=False, color_rgb=(15, 23, 42), align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        apply_run_font(run, font_name, size_pt, bold, italic, color_rgb)
    return p

def create_student_info_table(doc):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(3.4), Inches(3.4)]
    fields = [
        [("ชื่อ-นามสกุล: ", "................................................................"), ("รหัสนักศึกษา: ", "................................................................")],
        [("ระดับชั้น / สาขาวิชา: ", "สสส.2 / เทคโนโลยีสารสนเทศ"), ("วันที่ปฏิบัติงาน: ", "................................................................")]
    ]
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            set_cell_border(cell, top=dict(val='single', sz='4', color='CBD5E1'),
                                  bottom=dict(val='single', sz='4', color='CBD5E1'),
                                  left=dict(val='single', sz='4', color='CBD5E1'),
                                  right=dict(val='single', sz='4', color='CBD5E1'))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            label, dots = fields[row_idx][col_idx]
            r1 = p.add_run(label)
            apply_run_font(r1, bold=True, size_pt=14, color_rgb=(30, 41, 59))
            r2 = p.add_run(dots)
            apply_run_font(r2, bold=False, size_pt=14, color_rgb=(71, 85, 105))

def create_answer_box(doc, title="คำสั่งภาษา Python (พิมพ์โค้ดในกรอบนี้):", placeholder="df = ...\n\n", is_code=True, min_height_lines=4):
    p_title = add_styled_paragraph(doc, title, size_pt=14, bold=True, color_rgb=(30, 41, 59), space_before=4, space_after=2)
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.8)
    bg = "F1F5F9" if is_code else "FFFFFF"
    set_cell_background(cell, bg)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_border(cell, top=dict(val='single', sz='6', color='94A3B8'),
                          bottom=dict(val='single', sz='6', color='94A3B8'),
                          left=dict(val='single', sz='12', color='0284C7' if is_code else '10B981'),
                          right=dict(val='single', sz='6', color='94A3B8'))
    
    cell_p = cell.paragraphs[0]
    cell_p.paragraph_format.space_before = Pt(0)
    cell_p.paragraph_format.space_after = Pt(0)
    font_name = "Courier New" if is_code else FONT_NAME
    color_rgb = (51, 65, 85) if is_code else (71, 85, 105)
    r = cell_p.add_run(placeholder + ("\n" * (min_height_lines - 1)))
    apply_run_font(r, font_name=font_name, size_pt=14, color_rgb=color_rgb)

def generate_warmup_docx(output_path: Path):
    doc = docx.Document()
    
    # Page Margins 0.8 inch
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    # Header
    add_styled_paragraph(doc, "ใบงานทบทวนก่อนเรียน (Google Docs)", size_pt=20, bold=True, color_rgb=(14, 116, 144), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_styled_paragraph(doc, "รายวิชา 31901-2012 การวิเคราะห์และนำเสนอข้อมูล · สัปดาห์ที่ 2 วันจันทร์", size_pt=16, bold=True, color_rgb=(51, 65, 85), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_styled_paragraph(doc, "หัวข้อ: ทบทวนคำสั่งพื้นฐานการนำเข้าและสำรวจข้อมูลตาราง (เวลา 30 นาที)", size_pt=15, bold=False, italic=True, color_rgb=(71, 85, 105), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    
    create_student_info_table(doc)
    
    # Instructions Box
    inst_table = doc.add_table(rows=1, cols=1)
    inst_cell = inst_table.rows[0].cells[0]
    inst_cell.width = Inches(6.8)
    set_cell_background(inst_cell, "EFF6FF")
    set_cell_margins(inst_cell, top=100, bottom=100, left=160, right=160)
    set_cell_border(inst_cell, top=dict(val='single', sz='4', color='BFDBFE'),
                               bottom=dict(val='single', sz='4', color='BFDBFE'),
                               left=dict(val='single', sz='12', color='2563EB'),
                               right=dict(val='single', sz='4', color='BFDBFE'))
    inst_p = inst_cell.paragraphs[0]
    inst_p.paragraph_format.space_before = Pt(0)
    inst_p.paragraph_format.space_after = Pt(0)
    r_inst = inst_p.add_run(
        "คำชี้แจงสำหรับนักศึกษา:\n"
        "1. สามารถเปิดและพิมพ์ตอบในเอกสารนี้ผ่าน Google Docs หรือโปรแกรม Microsoft Word ได้โดยตรง\n"
        "2. ใช้ชุดข้อมูล nvc_tech_store.csv ที่อัปโหลดใน Google Colab ทำการทดลองรันคำสั่ง\n"
        "3. พิมพ์คำสั่งภาษา Python และพิมพ์สรุปผลลัพธ์ลงในกล่องข้อความใต้แต่ละข้อให้ครบถ้วน"
    )
    apply_run_font(r_inst, size_pt=14, color_rgb=(30, 58, 138))

    questions = [
        {
            "num": "1",
            "title": "ข้อที่ 1: ทบทวนการนำเข้าข้อมูลและการกำหนดชนิดข้อมูลหมายเลขโทรศัพท์",
            "prompt": "ให้ผู้เรียนเขียนคำสั่งนำเข้าไลบรารี pandas ในชื่อย่อ pd และอ่านไฟล์ nvc_tech_store.csv ลงในตัวแปรตาราง df โดยกำหนดให้คอลัมน์ customer_phone มีชนิดข้อมูลเป็นข้อความ (string) จากนั้นแสดงข้อมูลตัวอย่าง 3 แถวแรก",
            "funcs": "คำสั่งที่ทบทวน: import pandas as pd · pd.read_csv(..., dtype=...) · df.head(3)",
            "code_placeholder": "import pandas as pd\n# พิมพ์คำสั่งอ่านไฟล์และแสดงตัวอย่าง 3 แถวแรกตรงนี้\n",
            "ans_title": "สรุปผลลัพธ์จากการรันคำสั่ง (พิมพ์คำอธิบายและผลลัพธ์ในกรอบนี้):",
            "ans_placeholder": "[ ผลลัพธ์: ตารางมีขนาด 36 แถว 7 คอลัมน์ โดยคอลัมน์ customer_phone ยังคงแสดงเลข 08 นำหน้าครบถ้วน ]"
        },
        {
            "num": "2",
            "title": "ข้อที่ 2: ทบทวนการตรวจสอบขนาดของตารางและรายชื่อคอลัมน์",
            "prompt": "ให้ผู้เรียนเขียนคำสั่งตรวจสอบขนาดของตาราง (จำนวนแถวและจำนวนคอลัมน์) และแสดงรายชื่อคอลัมน์ทั้งหมดของตาราง df เพื่อตรวจสอบการสะกดชื่อคอลัมน์ก่อนนำไปคัดกรองข้อมูล",
            "funcs": "คำสั่งที่ทบทวน: df.shape · df.columns",
            "code_placeholder": "# พิมพ์คำสั่งตรวจสอบขนาดตารางและรายชื่อคอลัมน์ตรงนี้\n",
            "ans_title": "สรุปผลลัพธ์ขนาดของตารางและรายชื่อคอลัมน์ (พิมพ์ผลลัพธ์ในกรอบนี้):",
            "ans_placeholder": "[ ผลลัพธ์: ตารางมีขนาด (แถว, คอลัมน์) เท่าใด และมีคอลัมน์ใดบ้าง ]"
        },
        {
            "num": "3",
            "title": "ข้อที่ 3: ทบทวนการตรวจสอบข้อมูลสูญหายและค่าสถิติพื้นฐาน",
            "prompt": "ให้ผู้เรียนเขียนคำสั่งนับจำนวนข้อมูลสูญหายแยกตามรายคอลัมน์ และคำนวณค่าสถิติพื้นฐานของคอลัมน์ตัวเลข เพื่อระบุตำแหน่งปัญหาที่จะนำไปจัดการในกระบวนการ ETL วันนี้",
            "funcs": "คำสั่งที่ทบทวน: df.info() · df.isna().sum() · df.describe()",
            "code_placeholder": "# พิมพ์คำสั่งนับจำนวนค่าว่างและคำนวณค่าสถิติพื้นฐานตรงนี้\n",
            "ans_title": "สรุปผลการตรวจสอบข้อมูลสูญหายและค่าผิดปกติ (พิมพ์ผลลัพธ์ในกรอบนี้):",
            "ans_placeholder": "[ ผลลัพธ์: พบคอลัมน์ใดมีค่าว่างกี่จุด และพบคอลัมน์ตัวเลขใดมีค่าผิดปกติ (เช่น ติดลบ หรือสูงผิดปกติ) ]"
        }
    ]

    for q in questions:
        add_styled_paragraph(doc, "", space_before=8, space_after=0)
        add_styled_paragraph(doc, q["title"], size_pt=16, bold=True, color_rgb=(14, 116, 144), space_before=6, space_after=2)
        add_styled_paragraph(doc, q["funcs"], size_pt=13, bold=False, italic=True, color_rgb=(71, 85, 105), space_before=0, space_after=4)
        add_styled_paragraph(doc, f"โจทย์: {q['prompt']}", size_pt=15, bold=False, color_rgb=(15, 23, 42), space_before=2, space_after=6)
        
        create_answer_box(doc, "1) คำสั่งภาษา Python (พิมพ์โค้ดในกรอบนี้):", q["code_placeholder"], is_code=True, min_height_lines=3)
        create_answer_box(doc, f"2) {q['ans_title']}", q["ans_placeholder"], is_code=False, min_height_lines=2)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Generated {output_path}", flush=True)

def generate_full_monday_worksheet_docx(output_path: Path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    add_styled_paragraph(doc, "ใบงานปฏิบัติการประจำสัปดาห์ (Google Docs)", size_pt=20, bold=True, color_rgb=(14, 116, 144), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_styled_paragraph(doc, "รายวิชา 31901-2012 การวิเคราะห์และนำเสนอข้อมูล · สัปดาห์ที่ 2 วันจันทร์ (3 ชั่วโมง)", size_pt=16, bold=True, color_rgb=(51, 65, 85), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_styled_paragraph(doc, "ชุดข้อมูลสำหรับการฝึกปฏิบัติ: nvc_tech_store.csv (36 แถว 7 คอลัมน์)", size_pt=14, bold=False, italic=True, color_rgb=(71, 85, 105), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    
    create_student_info_table(doc)

    # Section 1
    add_styled_paragraph(doc, "ส่วนที่ 1: แบบฝึกหัดทบทวนก่อนเรียน (เวลา 30 นาที)", size_pt=18, bold=True, color_rgb=(3, 105, 161), space_before=14, space_after=4)
    w_questions = [
        ("ข้อที่ 1: ทบทวนการนำเข้าข้อมูลและกำหนดชนิดข้อมูลหมายเลขโทรศัพท์",
         "เขียนคำสั่งนำเข้าไลบรารี pandas ในชื่อย่อ pd และอ่านไฟล์ nvc_tech_store.csv ลงในตัวแปร df โดยกำหนดให้คอลัมน์ customer_phone มีชนิดข้อมูลเป็นข้อความ (string) จากนั้นแสดงข้อมูลตัวอย่าง 3 แถวแรก",
         "import pandas as pd\ndf = pd.read_csv('nvc_tech_store.csv', dtype={'customer_phone': 'string'})\ndf.head(3)"),
        ("ข้อที่ 2: ทบทวนการตรวจสอบขนาดตารางและรายชื่อคอลัมน์",
         "เขียนคำสั่งตรวจสอบขนาดของตารางด้วยคำสั่ง df.shape และแสดงรายชื่อคอลัมน์ทั้งหมดด้วยคำสั่ง df.columns",
         "df.shape\ndf.columns.tolist()"),
        ("ข้อที่ 3: ทบทวนการตรวจสอบข้อมูลสูญหายและค่าสถิติพื้นฐาน",
         "เขียนคำสั่งนับจำนวนค่าว่างแยกตามรายคอลัมน์ด้วยคำสั่ง df.isna().sum() และคำนวณค่าสถิติพื้นฐานด้วยคำสั่ง df.describe()",
         "df.isna().sum()\ndf.describe()")
    ]
    for idx, (t, p, c) in enumerate(w_questions, 1):
        add_styled_paragraph(doc, t, size_pt=15, bold=True, color_rgb=(30, 41, 59), space_before=6, space_after=2)
        add_styled_paragraph(doc, f"โจทย์: {p}", size_pt=14, bold=False, color_rgb=(51, 65, 85), space_before=0, space_after=4)
        create_answer_box(doc, "คำสั่งภาษา Python:", "# พิมพ์คำสั่งตรงนี้\n", is_code=True, min_height_lines=2)
        create_answer_box(doc, "สรุปผลลัพธ์:", "[ พิมพ์ผลลัพธ์ตรงนี้ ]", is_code=False, min_height_lines=2)

    # Section 2
    add_styled_paragraph(doc, "ส่วนที่ 2: แบบฝึกหัดปฏิบัติการท้ายคาบเรียน (เวลา 40 นาที · 10 คะแนน)", size_pt=18, bold=True, color_rgb=(3, 105, 161), space_before=16, space_after=4)
    add_styled_paragraph(doc, "คำชี้แจง: ให้ผู้เรียนใช้ตารางข้อมูลที่ผ่านการจัดเตรียมแล้ว (df_clean) ในการตอบคำถามทั้ง 5 ข้อ (ข้อละ 2 คะแนน)", size_pt=14, italic=True, color_rgb=(71, 85, 105), space_after=6)

    lab_questions = [
        ("ข้อที่ 1 (2 คะแนน): เปรียบเทียบขนาดของตารางข้อมูลก่อนและหลังทำ ETL",
         "เขียนคำสั่งแสดงขนาดของตาราง df.shape เปรียบเทียบกับ df_clean.shape และระบุว่ามีแถวข้อมูลที่มีปัญหาถูกคัดกรองออกไปทั้งหมดกี่แถว",
         "df.shape, df_clean.shape"),
        ("ข้อที่ 2 (2 คะแนน): ตรวจสอบจำนวนค่าว่างคงเหลือในตาราง df_clean",
         "เขียนคำสั่งตรวจสอบจำนวนค่าว่างคงเหลือด้วยคำสั่ง df_clean.isna().sum() และระบุว่าคอลัมน์ใดที่ยังคงมีค่าว่างเหลืออยู่ และเหลืออยู่กี่แถว พร้อมอธิบายเหตุผลที่ไม่ลบแถวดังกล่าวออก",
         "df_clean.isna().sum()"),
        ("ข้อที่ 3 (2 คะแนน): สรุปจำนวนรายการสินค้าแยกตามหมวดหมู่ภายหลังการจัดเตรียมข้อมูล",
         "เขียนคำสั่งนับจำนวนรายการสินค้าแต่ละหมวดหมู่ด้วยคำสั่ง df_clean['category'].value_counts() และระบุจำนวนรายการของทั้ง 3 หมวดหมู่",
         "df_clean['category'].value_counts()"),
        ("ข้อที่ 4 (2 คะแนน): ตรวจสอบค่าสถิติพื้นฐานของตาราง df_clean",
         "เขียนคำสั่งคำนวณค่าสถิติพื้นฐานด้วยคำสั่ง df_clean.describe() และระบุค่าเฉลี่ย (mean) ค่าต่ำสุด (min) และค่าสูงสุด (max) ของคอลัมน์ total_price ภายหลังการจัดเตรียมข้อมูล",
         "df_clean.describe()"),
        ("ข้อที่ 5 (2 คะแนน): ตรวจสอบความถูกต้องของไฟล์ CSV ที่บันทึกออก",
         "เขียนคำสั่งอ่านไฟล์ nvc_tech_store_clean.csv ขึ้นมาแสดงข้อมูลตัวอย่าง 5 แถวแรกด้วยคำสั่ง pd.read_csv('nvc_tech_store_clean.csv').head() เพื่อยืนยันความถูกต้องของไฟล์",
         "pd.read_csv('nvc_tech_store_clean.csv').head()")
    ]
    for idx, (t, p, c) in enumerate(lab_questions, 1):
        add_styled_paragraph(doc, t, size_pt=15, bold=True, color_rgb=(30, 41, 59), space_before=6, space_after=2)
        add_styled_paragraph(doc, f"โจทย์: {p}", size_pt=14, bold=False, color_rgb=(51, 65, 85), space_before=0, space_after=4)
        create_answer_box(doc, "คำสั่งภาษา Python:", "# พิมพ์คำสั่งตรงนี้\n", is_code=True, min_height_lines=2)
        create_answer_box(doc, "สรุปคำตอบจากผลลัพธ์ (2 คะแนน):", "[ พิมพ์คำตอบตรงนี้ ]", is_code=False, min_height_lines=2)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Generated {output_path}", flush=True)

def generate_teacher_answer_key_docx(output_path: Path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    add_styled_paragraph(doc, "เฉลยใบงานปฏิบัติการประจำสัปดาห์ (สำหรับผู้สอน)", size_pt=20, bold=True, color_rgb=(185, 28, 28), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_styled_paragraph(doc, "รายวิชา 31901-2012 การวิเคราะห์และนำเสนอข้อมูล · สัปดาห์ที่ 2 วันจันทร์", size_pt=16, bold=True, color_rgb=(51, 65, 85), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_styled_paragraph(doc, "เอกสารเฉลยและเกณฑ์การให้คะแนนแบบฝึกหัดทบทวนและแบบฝึกหัดท้ายคาบ", size_pt=14, italic=True, color_rgb=(100, 116, 139), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    add_styled_paragraph(doc, "ส่วนที่ 1: เฉลยแบบฝึกหัดทบทวนก่อนเรียน (เวลา 30 นาที)", size_pt=18, bold=True, color_rgb=(185, 28, 28), space_before=10, space_after=6)
    w_keys = [
        ("ข้อที่ 1: ทบทวนการนำเข้าข้อมูลและกำหนดชนิดข้อมูลหมายเลขโทรศัพท์",
         "import pandas as pd\ndf = pd.read_csv('nvc_tech_store.csv', dtype={'customer_phone': 'string'})\ndf.head(3)",
         "นำเข้าตารางข้อมูลต้นฉบับขนาด 36 แถว 7 คอลัมน์ โดยคอลัมน์ customer_phone ยังคงแสดงเลข 08 นำหน้าครบถ้วน"),
        ("ข้อที่ 2: ทบทวนการตรวจสอบขนาดตารางและรายชื่อคอลัมน์",
         "df.shape, df.columns.tolist()",
         "ตารางข้อมูลต้นฉบับมีขนาด (36, 7) คือ 36 แถว และ 7 คอลัมน์ รายชื่อคอลัมน์ได้แก่ order_id, customer_phone, product_name, category, quantity, total_price, payment_status"),
        ("ข้อที่ 3: ทบทวนการตรวจสอบข้อมูลสูญหายและค่าสถิติพื้นฐาน",
         "df.isna().sum()\ndf.describe()",
         "พบค่าว่างรวม 5 จุด (customer_phone 2, quantity 1, total_price 2 แถว) และพบคอลัมน์ quantity มีค่าต่ำสุด -2 ชิ้น ส่วน total_price มีค่าสูงสุด 99,000 บาท ส่งผลให้ค่าเฉลี่ยคลาดเคลื่อนไปอยู่ที่ 3,613.82 บาท")
    ]
    for t, c, a in w_keys:
        add_styled_paragraph(doc, t, size_pt=15, bold=True, color_rgb=(30, 41, 59), space_before=4, space_after=2)
        create_answer_box(doc, "คำสั่งเฉลยภาษา Python:", c, is_code=True, min_height_lines=2)
        create_answer_box(doc, "แนวคำตอบและคำอธิบายผลลัพธ์:", a, is_code=False, min_height_lines=2)

    add_styled_paragraph(doc, "ส่วนที่ 2: เฉลยแบบฝึกหัดท้ายคาบเรียนวันจันทร์ (10 คะแนน)", size_pt=18, bold=True, color_rgb=(185, 28, 28), space_before=16, space_after=6)
    lab_keys = [
        ("ข้อที่ 1 (2 คะแนน): เปรียบเทียบขนาดของตารางข้อมูลก่อนและหลังทำ ETL",
         "df.shape, df_clean.shape",
         "ตารางข้อมูลต้นฉบับมี 36 แถว และตารางข้อมูลที่ผ่านการจัดเตรียมแล้วเหลือ 29 แถว แสดงว่ามีแถวข้อมูลที่มีปัญหาถูกคัดกรองออกไปทั้งหมด 7 แถว"),
        ("ข้อที่ 2 (2 คะแนน): ตรวจสอบจำนวนค่าว่างคงเหลือในตาราง df_clean",
         "df_clean.isna().sum()",
         "คอลัมน์ customer_phone ยังคงมีค่าว่างเหลืออยู่ 2 แถว เนื่องจากเป็นรายการสั่งซื้อที่มีมูลค่าการซื้อขายครบถ้วนสมบูรณ์ จึงเก็บรักษาไว้คำนวณยอดขาย"),
        ("ข้อที่ 3 (2 คะแนน): สรุปจำนวนรายการสินค้าแยกตามหมวดหมู่ภายหลังการจัดเตรียมข้อมูล",
         "df_clean['category'].value_counts()",
         "หมวดหมู่อุปกรณ์ต่อพ่วง มี 12 รายการ, สายและหน่วยความจำ มี 11 รายการ และอุปกรณ์เสียง มี 6 รายการ รวมทั้งสิ้น 29 รายการ ครบ 3 หมวดหมู่มาตรฐาน"),
        ("ข้อที่ 4 (2 คะแนน): ตรวจสอบค่าสถิติพื้นฐานของตาราง df_clean",
         "df_clean.describe()",
         "ภายหลังการจัดเตรียมข้อมูล คอลัมน์ total_price มีค่าเฉลี่ย (mean) เท่ากับ 754.83 บาท, ค่าต่ำสุด (min) เท่ากับ 180.00 บาท และค่าสูงสุด (max) เท่ากับ 2,580.00 บาท"),
        ("ข้อที่ 5 (2 คะแนน): ตรวจสอบความถูกต้องของไฟล์ CSV ที่บันทึกออก",
         "pd.read_csv('nvc_tech_store_clean.csv').head()",
         "ไฟล์ nvc_tech_store_clean.csv สามารถอ่านกลับมาใช้งานได้ถูกต้อง มี 7 คอลัมน์มาตรฐาน โดยไม่มีคอลัมน์ Unnamed: 0 เพิ่มขึ้นมา")
    ]
    for t, c, a in lab_keys:
        add_styled_paragraph(doc, t, size_pt=15, bold=True, color_rgb=(30, 41, 59), space_before=4, space_after=2)
        create_answer_box(doc, "คำสั่งเฉลยภาษา Python:", c, is_code=True, min_height_lines=2)
        create_answer_box(doc, "แนวคำตอบ (2 คะแนน):", a, is_code=False, min_height_lines=2)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Generated {output_path}", flush=True)

if __name__ == "__main__":
    base = Path("/Users/surachartlimrattanaphun/Desktop/Teaching_NVC")
    dva = base / "วิชาการวิเคราะห์และนำเสนอข้อมูล"
    w2_retro = dva / "ชุดพร้อมสอน_31901-2012/เรียนย้อนหลัง_สัปดาห์02"
    public_w2 = dva / "dva-slides/public/ชุดเรียน/เรียนย้อนหลัง_สัปดาห์02"
    teacher_dir = dva / "ชุดพร้อมสอน_31901-2012/สำหรับครู"
    deliverables = base / "deliverables"

    for target_dir in [w2_retro, public_w2, deliverables]:
        generate_warmup_docx(target_dir / "ใบงานทบทวนก่อนเรียน_สัปดาห์ที่02_GoogleDocs.docx")
        generate_full_monday_worksheet_docx(target_dir / "ใบงานปฏิบัติการ_สัปดาห์ที่02_จันทร์_GoogleDocs.docx")

    generate_teacher_answer_key_docx(teacher_dir / "เฉลยใบงานปฏิบัติการ_สัปดาห์ที่02_จันทร์_สำหรับผู้สอน.docx")
    generate_teacher_answer_key_docx(deliverables / "เฉลยใบงานปฏิบัติการ_สัปดาห์ที่02_จันทร์_สำหรับผู้สอน.docx")
    print("All Google Docs worksheets generated successfully.")
